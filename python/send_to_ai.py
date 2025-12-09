import os
import re
import json
import ollama
import zipfile
from docx import Document


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def extract_text(json_data):
    """Extrait le texte selon la structure du JSON (paragraphs ou structure)"""
    if "paragraphs" in json_data:
        return "\n".join(json_data["paragraphs"])
    elif "structure" in json_data:
        # Nouvelle structure hiérarchique
        texts = []
        for section in json_data.get("structure", []):
            texts.append(extract_from_section(section))
        return "\n".join(texts)
    elif "context_files" in json_data:
        texts = []
        for ctx in json_data["context_files"]:
            texts.append(extract_text(ctx))
        return "\n".join(texts)
    return ""


def extract_from_section(section):
    """Extrait récursivement le texte d'une section"""
    parts = []
    if "title" in section:
        parts.append(section["title"])
    
    # Contenus directs de la section
    for content in section.get("contents", []):
        if isinstance(content, dict) and content.get("type") == "placeholder":
            parts.append(content.get("full_text", ""))
        elif isinstance(content, str):
            parts.append(content)
    
    # Sous-sections
    for subsection in section.get("subsections", []):
        parts.append(extract_from_section(subsection))
    
    # Sous-sous-sections
    for subsubsection in section.get("subsubsections", []):
        parts.append(extract_from_section(subsubsection))
    
    return "\n".join(parts)


def extract_placeholders_from_section(section, placeholders, path_prefix=""):
    """Extrait récursivement tous les placeholders <here>...</here> d'une section"""
    section_title = section.get("title", "untitled")
    current_path = f"{path_prefix}/{section_title}" if path_prefix else section_title
    
    # Extraire les placeholders des contenus directs
    for idx, content in enumerate(section.get("contents", [])):
        if isinstance(content, dict) and content.get("type") == "placeholder":
            full_text = content.get("full_text", "")
            for occ_idx, match in enumerate(re.finditer(r"<here>(.*?)</here>", full_text, flags=re.S)):
                original = match.group(1).strip()
                pid = f"s{len(placeholders)}_h{occ_idx}"
                placeholders.append({
                    "id": pid,
                    "section_path": current_path,
                    "content_index": idx,
                    "occurrence_index": occ_idx,
                    "original": original,
                    "full_text": full_text
                })
    
    # Traiter les sous-sections
    for subsection in section.get("subsections", []):
        extract_placeholders_from_section(subsection, placeholders, current_path)
    
    # Traiter les sous-sous-sections
    for subsubsection in section.get("subsubsections", []):
        extract_placeholders_from_section(subsubsection, placeholders, current_path)





def complete_text_with_ai(context_text, placeholders_batch, model_name="llama3.1:8b"):
    """
    Traite un lot de placeholders à la fois avec Ollama
    """
    # Construire la liste des placeholders pour le prompt
    placeholders_text = ""
    for i, ph in enumerate(placeholders_batch):
        section_path = ph.get("section_path", "N/A")
        original = ph.get("original", "")
        placeholders_text += f"\n[{i+1}] Section: {section_path}\nOriginal: {original[:200]}...\n"

    prompt = f"""Vous êtes un assistant de documentation technique. Votre rôle est de générer le contenu des sections de documentation en vous basant strictement sur les informations extraites des documents fournis (cadrage, spécifications, manuel, fichiers d'environnement).

CONSIGNES IMPORTANTES :
- Vous devez fournir EXACTEMENT {len(placeholders_batch)} remplacements (un pour chaque zone marquée par <here></here>)
- Le style doit être concis, structuré et professionnel
- Rédigez dans la langue majoritaire dans le contexte
- Basez-vous fortement sur les informations présentes dans le contexte fourni afin d'appuyer et d'etayer vos propositions
- N'inventez PAS d'informations non présentes dans le contexte
- Chaque remplacement doit faire 3 à 5 phrases

FORMAT DE SORTIE REQUIS :
Retournez un objet JSON VALIDE UNIQUEMENT (sans commentaires, sans texte supplémentaire) avec cette structure :
{{
  "replacements": [
    {{
      "index": 0,
      "replacement": "Texte de remplacement basé sur le contexte"
    }},
    {{
      "index": 1,
      "replacement": "Texte de remplacement basé sur le contexte"
    }}
    ... (continuez pour tous les {len(placeholders_batch)} emplacements)
  ]
}}

EXIGENCES CRITIQUES :
- Sortie JSON VALIDE uniquement (pas de markdown, pas de commentaires, pas de texte supplémentaire)
- Vous DEVEZ fournir les {len(placeholders_batch)} remplacements (index 0 à {len(placeholders_batch)-1})
- NE sautez AUCUN index
- N'ajoutez PAS de commentaires comme // ou /* */
- Basez chaque réponse sur le contexte fourni ci-dessous

CONTEXTE (documents de référence) :
{context_text[:5000]}

ZONES À COMPLÉTER (VOUS DEVEZ TOUTES LES TRAITER - {len(placeholders_batch)} au total) :
{placeholders_text}

RAPPEL : SORTIE JSON VALIDE AVEC EXACTEMENT {len(placeholders_batch)} REMPLACEMENTS, AUCUN COMMENTAIRE.
"""

    try:
        # Appel à Ollama
        print(f"   [DEBUG] Envoi de {len(placeholders_batch)} placeholders a Ollama...")
        response = ollama.chat(
            model=model_name,
            messages=[
                {
                    'role': 'user',
                    'content': prompt
                }
            ]
        )
        
        content = response['message']['content'].strip()
        print(f"   [DEBUG] Reponse brute recue ({len(content)} caracteres)")
        
    except Exception as e:
        print(f"   [ERROR] Exception Ollama: {str(e)}")
        return {"error": f"Erreur lors de la generation: {str(e)}"}

    # Parser la réponse JSON
    try:
        # Supprimer les balises markdown si présentes
        if content.startswith('```'):
            lines = content.split('\n')
            content = '\n'.join(lines[1:-1]) if len(lines) > 2 else content
            print("   [DEBUG] Balises markdown supprimees")
        
        parsed = json.loads(content)
        print(f"   [DEBUG] JSON parse avec succes")
    except Exception as e:
        print(f"   [WARN] Erreur parsing JSON direct: {str(e)}")
        # Nettoyer le contenu avant de le parser
        content_clean = re.sub(r'//.*?$', '', content, flags=re.MULTILINE)
        content_clean = re.sub(r'/\*.*?\*/', '', content_clean, flags=re.DOTALL)
        
        try:
            parsed = json.loads(content_clean)
            print(f"   [DEBUG] JSON parse apres nettoyage")
        except Exception:
            print(f"   [WARN] Tentative d'extraction du JSON...")
            m = re.search(r"\{.*\}", content_clean, flags=re.S)
            if m:
                try:
                    parsed = json.loads(m.group(0))
                    print(f"   [DEBUG] JSON extrait et parse")
                except Exception:
                    print(f"   [ERROR] Impossible de parser le JSON extrait")
                    return {"error": "Could not parse model response as JSON", "raw": content}
            else:
                print(f"   [ERROR] Aucun JSON trouve dans la reponse")
                return {"error": "No JSON found in model response", "raw": content}
    
    # Vérifier le nombre de remplacements
    replacements = parsed.get("replacements", [])
    print(f"   [DEBUG] Nombre de remplacements recus: {len(replacements)}/{len(placeholders_batch)}")
    
    if len(replacements) < len(placeholders_batch):
        print(f"   [WARN] Ollama n'a retourne que {len(replacements)} remplacements sur {len(placeholders_batch)} demandes")

    return parsed



def create_completed_docx(original_docx_path, replacement_pairs, output_path):
    """
    Crée un nouveau document Word en remplaçant les placeholders <here>...</here>
    par les propositions du LLM
    """
    print(f"\n[INFO] Creation du document Word complete...")
    
    # Charger le document original
    doc = Document(original_docx_path)
    
    # Créer un dictionnaire pour accès rapide aux remplacements
    replacements_map = {}
    for pair in replacement_pairs:
        original_with_tags = pair.get("original_text_with_tags", "")
        replacement = pair.get("proposed_replacement", "")
        # Extraire le contenu entre <here> et </here>
        match = re.search(r"<here>(.*?)</here>", original_with_tags, flags=re.S)
        if match:
            original_content = match.group(1).strip()
            replacements_map[original_content] = replacement
    
    print(f"   [DEBUG] {len(replacements_map)} remplacements a appliquer")
    
    # Parcourir tous les paragraphes et remplacer
    replacements_done = 0
    for para in doc.paragraphs:
        text = para.text
        
        # Vérifier si le paragraphe contient un placeholder
        if '<here>' in text and '</here>' in text:
            # Extraire le contenu entre les balises
            match = re.search(r"<here>(.*?)</here>", text, flags=re.S)
            if match:
                original_content = match.group(1).strip()
                
                # Chercher le remplacement correspondant
                if original_content in replacements_map:
                    replacement_text = replacements_map[original_content]
                    
                    # Remplacer tout le contenu du paragraphe
                    para.text = replacement_text
                    replacements_done += 1
                    print(f"   [OK] Remplacement {replacements_done} applique")
    
    # Sauvegarder le document complété
    doc.save(output_path)
    print(f"   [OK] Document Word complete sauvegarde: {output_path}")
    print(f"   [OK] {replacements_done}/{len(replacements_map)} remplacements appliques")
    
    return replacements_done


def create_zip_package(json_path, docx_path, zip_output_path):
    """
    Crée un fichier ZIP contenant le JSON et le DOCX
    """
    print(f"\n[INFO] Creation du package ZIP...")
    
    with zipfile.ZipFile(zip_output_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
        # Ajouter le JSON
        zipf.write(json_path, os.path.basename(json_path))
        print(f"   [OK] Ajoute: {os.path.basename(json_path)}")
        
        # Ajouter le DOCX
        zipf.write(docx_path, os.path.basename(docx_path))
        print(f"   [OK] Ajoute: {os.path.basename(docx_path)}")
    
    print(f"   [OK] Package ZIP cree: {zip_output_path}")
    return zip_output_path



def main():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    root_dir = os.path.dirname(base_dir)
    json_dir = os.path.join(root_dir, "jsons")

    context_json = os.path.join(json_dir, "context_global.json")
    principal_json = os.path.join(json_dir, "principal.json")

    context_data = load_json(context_json)
    principal_data = load_json(principal_json)

    context_text = extract_text(context_data)
    document_text = extract_text(principal_data)

    # Identify all <here>...</here> placeholders
    placeholders = []
    
    # Si c'est une structure hiérarchique
    if "structure" in principal_data:
        for section in principal_data.get("structure", []):
            extract_placeholders_from_section(section, placeholders)
    # Sinon, ancienne méthode avec paragraphs
    elif "paragraphs" in principal_data:
        paragraphs = principal_data.get("paragraphs", [])
        for idx, para in enumerate(paragraphs):
            for occ_idx, match in enumerate(re.finditer(r"<here>(.*?)</here>", para, flags=re.S)):
                original = match.group(1).strip()
                pid = f"p{idx}_h{occ_idx}"
                placeholders.append({
                    "id": pid,
                    "paragraph_index": idx,
                    "occurrence_index": occ_idx,
                    "original": original,
                    "context_path": f"paragraph[{idx}]"
                })

    print("[INFO] Envoi du contexte et du document a l'IA...")
    print(f"   [INFO] {len(placeholders)} zones <here>...</here> detectees")
    
    # Traiter par lots de 3 placeholders à la fois (réduit pour améliorer la fiabilité)
    BATCH_SIZE = 3
    all_replacement_pairs = []
    total_batches = (len(placeholders) + BATCH_SIZE - 1) // BATCH_SIZE
    
    for batch_num in range(total_batches):
        start_idx = batch_num * BATCH_SIZE
        end_idx = min(start_idx + BATCH_SIZE, len(placeholders))
        batch = placeholders[start_idx:end_idx]
        
        print(f"   [BATCH {batch_num + 1}/{total_batches}] Traitement de {len(batch)} zones...")
        
        model_response = complete_text_with_ai(context_text, batch)
        
        if isinstance(model_response, dict) and model_response.get("error"):
            print(f"   [WARN] Erreur pour le lot {batch_num + 1}: {model_response.get('error')}")
            continue
        
        parsed = model_response if isinstance(model_response, dict) else {}
        replacements = parsed.get("replacements") if isinstance(parsed, dict) else []
        if not isinstance(replacements, list):
            replacements = []
        
        # Créer les paires pour ce batch
        for rep in replacements:
            rep_index = rep.get("index")
            repl_text = rep.get("replacement")
            
            if rep_index is None or repl_text is None:
                continue
            
            # Récupérer le placeholder correspondant
            if 0 <= rep_index < len(batch):
                ph = batch[rep_index]
                pair = {
                    "original_text_with_tags": f"<here>{ph.get('original', '')}</here>",
                    "proposed_replacement": repl_text
                }
                
                if "section_path" in ph:
                    pair["section_path"] = ph["section_path"]
                    pair["content_index"] = ph.get("content_index")
                elif "paragraph_index" in ph:
                    pair["paragraph_index"] = ph["paragraph_index"]
                
                all_replacement_pairs.append(pair)
        
        print(f"   [OK] Lot {batch_num + 1} traite : {len(replacements)} remplacements generes")

    if not all_replacement_pairs and total_batches > 0:
        error_result = {
            "error": "Aucun remplacement n'a pu etre genere",
            "total_placeholders": len(placeholders),
            "total_batches_attempted": total_batches
        }
        output_json_path = os.path.join(json_dir, "completed_document.json")
        with open(output_json_path, "w", encoding="utf-8") as f:
            json.dump(error_result, f, indent=2, ensure_ascii=False)
        print(f"[ERROR] Aucun remplacement genere")
        return

    result_json = {
        "file": principal_data.get("file", "document.docx"),
        "summary": {
            "total_placeholders_found": len(placeholders),
            "total_replacements_proposed": len(all_replacement_pairs),
            "batches_processed": total_batches
        },
        "replacement_pairs": all_replacement_pairs
    }

    output_json_path = os.path.join(json_dir, "completed_document.json")
    with open(output_json_path, "w", encoding="utf-8") as f:
        json.dump(result_json, f, indent=2, ensure_ascii=False)

    print(f"\n[OK] JSON genere: {output_json_path}")
    print(f"   [OK] {len(all_replacement_pairs)}/{len(placeholders)} remplacements proposes")
    
    # Générer le document Word complété
    tocomplete_dir = os.path.join(root_dir, "tocomplete")
    original_docx = None
    
    # Trouver le fichier .docx original
    if os.path.exists(tocomplete_dir):
        for f in os.listdir(tocomplete_dir):
            if f.endswith(".docx") and not f.startswith("~$"):
                original_docx = os.path.join(tocomplete_dir, f)
                break
    
    if original_docx and os.path.exists(original_docx):
        output_docx_path = os.path.join(json_dir, "completed_document.docx")
        
        try:
            replacements_done = create_completed_docx(original_docx, all_replacement_pairs, output_docx_path)
            
            # Créer le package ZIP
            zip_output_path = os.path.join(json_dir, "completion_results.zip")
            create_zip_package(output_json_path, output_docx_path, zip_output_path)
            
            print(f"\n[OK] Traitement termine avec succes !")
            print(f"   - JSON: {output_json_path}")
            print(f"   - DOCX: {output_docx_path}")
            print(f"   - ZIP: {zip_output_path}")
            
        except Exception as e:
            print(f"\n[WARN] Erreur lors de la creation du DOCX: {str(e)}")
            print(f"   [INFO] Le JSON a ete genere: {output_json_path}")
    else:
        print(f"\n[WARN] Document original non trouve, seul le JSON a ete genere")
        print(f"   - JSON: {output_json_path}")


if __name__ == "__main__":
    main()
