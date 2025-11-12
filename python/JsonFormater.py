import os
import json
import sys
import shutil
import re
from docx import Document

# --- Fonctions utilitaires ---
def extract_structured_content_from_docx(path):
    """
    Extrait uniquement les titres (headings) et les contenus avec <here>...</here>
    Retourne une structure hiérarchique avec sections, sous-sections et contenus
    
    Gère les cas où <here> et </here> sont dans des paragraphes différents
    """
    doc = Document(path)
    structure = []
    current_section = None
    current_subsection = None
    current_subsubsection = None
    
    # Variables pour gérer les blocs <here>...</here> multi-paragraphes
    inside_here_block = False
    here_content = []
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        
        # Déterminer le niveau de titre basé sur le style
        style_name = para.style.name if para.style else ""
        
        # Si on est dans un bloc <here>, accumuler le contenu
        if inside_here_block:
            here_content.append(text)
            # Vérifier si on ferme le bloc
            if '</here>' in text:
                inside_here_block = False
                full_text = '\n'.join(here_content)
                content_item = {
                    "type": "placeholder",
                    "full_text": full_text
                }
                
                # Ajouter au conteneur approprié
                if current_subsubsection and "contents" in current_subsubsection:
                    current_subsubsection["contents"].append(content_item)
                elif current_subsection and "contents" in current_subsection:
                    current_subsection["contents"].append(content_item)
                elif current_section:
                    if "contents" not in current_section:
                        current_section["contents"] = []
                    current_section["contents"].append(content_item)
                else:
                    # Contenu orphelin
                    if not structure or "contents" not in structure[-1]:
                        structure.append({
                            "title": "Contenu sans titre",
                            "level": 0,
                            "contents": []
                        })
                    structure[-1]["contents"].append(content_item)
                
                here_content = []
            continue
        
        # Vérifier si on ouvre un bloc <here>
        if '<here>' in text:
            inside_here_block = True
            here_content = [text]
            # Cas où <here> et </here> sont sur la même ligne
            if '</here>' in text:
                inside_here_block = False
                content_item = {
                    "type": "placeholder",
                    "full_text": text
                }
                
                # Ajouter au conteneur approprié
                if current_subsubsection and "contents" in current_subsubsection:
                    current_subsubsection["contents"].append(content_item)
                elif current_subsection and "contents" in current_subsection:
                    current_subsection["contents"].append(content_item)
                elif current_section:
                    if "contents" not in current_section:
                        current_section["contents"] = []
                    current_section["contents"].append(content_item)
                else:
                    # Contenu orphelin
                    if not structure or "contents" not in structure[-1]:
                        structure.append({
                            "title": "Contenu sans titre",
                            "level": 0,
                            "contents": []
                        })
                    structure[-1]["contents"].append(content_item)
                
                here_content = []
            continue
        
        # Vérifier si c'est un titre de niveau 1
        if style_name.startswith('Heading 1') or style_name == 'Title':
            current_section = {
                "title": text,
                "level": 1,
                "subsections": []
            }
            structure.append(current_section)
            current_subsection = None
            current_subsubsection = None
            
        # Vérifier si c'est un titre de niveau 2
        elif style_name.startswith('Heading 2'):
            if current_section:
                current_subsection = {
                    "title": text,
                    "level": 2,
                    "contents": []
                }
                current_section["subsections"].append(current_subsection)
            else:
                # Si pas de section parente, créer une section
                current_section = {
                    "title": text,
                    "level": 2,
                    "subsections": []
                }
                structure.append(current_section)
                current_subsection = None
            current_subsubsection = None
                
        # Vérifier si c'est un titre de niveau 3+
        elif style_name.startswith('Heading'):
            level = 3
            if 'Heading 3' in style_name:
                level = 3
            elif 'Heading 4' in style_name:
                level = 4
            elif 'Heading 5' in style_name:
                level = 5
                
            sub_item = {
                "title": text,
                "level": level,
                "contents": []
            }
            
            if current_subsection:
                # Ajouter comme sous-sous-section
                if "subsubsections" not in current_subsection:
                    current_subsection["subsubsections"] = []
                current_subsection["subsubsections"].append(sub_item)
                current_subsubsection = sub_item
            elif current_section:
                current_subsection = sub_item
                current_section["subsections"].append(current_subsection)
                current_subsubsection = None
    
    return structure


def extract_text_from_docx(path):
    """Version simple pour les fichiers de contexte"""
    doc = Document(path)
    content = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            content.append(text)
    return content


def convert_docx_to_json(input_path, output_path, use_structure=False):
    """
    Convertit un docx en JSON
    Si use_structure=True, extrait uniquement headings et <here>...</here>
    Sinon, extrait tout le texte
    """
    if use_structure:
        content = extract_structured_content_from_docx(input_path)
        data = {
            "file": os.path.basename(input_path), 
            "structure": content
        }
    else:
        content = extract_text_from_docx(input_path)
        data = {
            "file": os.path.basename(input_path), 
            "paragraphs": content
        }
    
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    print(f"[OK] {input_path} → {output_path}")



def merge_json_files(input_files, output_file):
    merged = {"context_files": []}
    for fpath in input_files:
        with open(fpath, "r", encoding="utf-8") as f:
            merged["context_files"].append(json.load(f))
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(merged, f, indent=2, ensure_ascii=False)
    print(f"[OK] Contexte global créé : {output_file}")


# --- Programme principal ---
def main():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    root_dir = os.path.dirname(base_dir)

    context_dir = os.path.join(root_dir, "context")
    json_dir = os.path.join(root_dir, "jsons")
    
    # Vider le contenu du dossier jsons (mais pas le dossier lui-même car c'est un volume monté)
    if os.path.exists(json_dir):
        for filename in os.listdir(json_dir):
            file_path = os.path.join(json_dir, filename)
            try:
                if os.path.isfile(file_path) or os.path.islink(file_path):
                    os.unlink(file_path)
                elif os.path.isdir(file_path):
                    shutil.rmtree(file_path)
            except Exception as e:
                print(f'⚠️  Erreur lors de la suppression de {file_path}: {e}')
        print(f"🗑️  Contenu du dossier jsons vidé")
    
    # S'assurer que le dossier existe
    os.makedirs(json_dir, exist_ok=True)

    # 1️⃣ Transformer les fichiers du dossier context/
    context_jsons = []
    for fname in os.listdir(context_dir):
        # Ignorer les fichiers temporaires Word (commencent par ~$)
        if fname.endswith(".docx") and not fname.startswith("~$"):
            src = os.path.join(context_dir, fname)
            dst = os.path.join(json_dir, f"{os.path.splitext(fname)[0]}.json")
            convert_docx_to_json(src, dst)
            context_jsons.append(dst)

    # Fusionner en un seul JSON global
    context_global_path = os.path.join(json_dir, "context_global.json")
    merge_json_files(context_jsons, context_global_path)

    # 2️⃣ Transformer le fichier principal à la racine avec structure
    main_docx = None
    for f in os.listdir(root_dir):
        # Ignorer les fichiers temporaires Word (commencent par ~$)
        if f.endswith(".docx") and not f.startswith("~$"):
            main_docx = os.path.join(root_dir, f)
            break

    if not main_docx:
        print("[ERREUR] Aucun fichier principal .docx trouvé à la racine.")
        sys.exit(1)

    principal_json = os.path.join(json_dir, "principal.json")
    # Utiliser use_structure=True pour extraire uniquement headings et <here>...</here>
    convert_docx_to_json(main_docx, principal_json, use_structure=True)

    print(f"\n✅ Conversion terminée.\n- Contexte global : {context_global_path}\n- Principal : {principal_json}")


if __name__ == "__main__":
    main()
